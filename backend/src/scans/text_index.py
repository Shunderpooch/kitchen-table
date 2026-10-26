"""Searchable text of every filed document, in the shared SQLite database.

``scanned_documents`` holds one row per filed PDF: where it is, what its
filename says (date, time, merchant, person), and what was read from it.
``scanned_pages`` holds the text of each page, and ``scanned_pages_fts`` is
an SQLite FTS5 index over it for full-text search. ``export_xlsx`` writes
both tables to Excel.

This is the permanent record of what was read, so OCR runs once per file
content: unchanged files are skipped, and a file that was renamed, moved or
re-dated is recognised by its sha256 and keeps its text. Each page records
how its text was obtained (``pdf_text``, ``ocr:<engine>``). Indexing only
reads files; online-only OneDrive files are skipped unless asked for,
because reading them downloads them.
"""

from __future__ import annotations

import hashlib
import os
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Callable, Iterable, Optional

from sqlalchemy import Engine, select, text
from sqlalchemy.orm import Session

from ..db.models import ScannedDocument, ScannedPage
from .conventions import OneDriveLayout
from .cutter import PageRef, page_text_with_source
from .extract import find_date_and_time, find_total
from .naming import parse_filename

EXCEL_CELL_LIMIT = 32_000  # Excel allows 32,767 characters per cell

_FTS_SQL = [
    """CREATE VIRTUAL TABLE IF NOT EXISTS scanned_pages_fts
       USING fts5(text, content='scanned_pages', content_rowid='id')""",
    """CREATE TRIGGER IF NOT EXISTS scanned_pages_ai AFTER INSERT ON scanned_pages BEGIN
         INSERT INTO scanned_pages_fts(rowid, text) VALUES (new.id, new.text);
       END""",
    """CREATE TRIGGER IF NOT EXISTS scanned_pages_ad AFTER DELETE ON scanned_pages BEGIN
         INSERT INTO scanned_pages_fts(scanned_pages_fts, rowid, text) VALUES ('delete', old.id, old.text);
       END""",
    """CREATE TRIGGER IF NOT EXISTS scanned_pages_au AFTER UPDATE ON scanned_pages BEGIN
         INSERT INTO scanned_pages_fts(scanned_pages_fts, rowid, text) VALUES ('delete', old.id, old.text);
         INSERT INTO scanned_pages_fts(rowid, text) VALUES (new.id, new.text);
       END""",
]


def ensure_schema(engine: Engine) -> None:
    ScannedDocument.__table__.create(engine, checkfirst=True)
    ScannedPage.__table__.create(engine, checkfirst=True)
    with engine.begin() as connection:
        columns = {row[1] for row in connection.exec_driver_sql("PRAGMA table_info(scanned_documents)")}
        if "sha256" not in columns:
            connection.exec_driver_sql("ALTER TABLE scanned_documents ADD COLUMN sha256 VARCHAR(64)")
            connection.exec_driver_sql(
                "CREATE INDEX IF NOT EXISTS ix_scanned_documents_sha256 ON scanned_documents (sha256)"
            )
        for statement in _FTS_SQL:
            connection.exec_driver_sql(statement)


def default_roots(layout: OneDriveLayout) -> list[Path]:
    documents = layout.documents
    roots = [layout.receipts, documents / "cars", documents / "manuals", documents / "hardware_instructions"]
    roots += [p.receipts_root for p in layout.people.values() if p.receipts_root_exists]
    return [r for r in roots if r.exists()]


def _kind_for(rel_parts: list[str]) -> str:
    lowered = [p.lower() for p in rel_parts]
    if "manuals" in lowered:
        return "manual"
    if "hardware_instructions" in lowered:
        return "instructions"
    if "cars" in lowered:
        return "car"
    return "receipt"


def _parse_date(value: Optional[str]) -> Optional[date]:
    if not value:
        return None
    try:
        return date.fromisoformat(value if len(value) == 10 else value + "-01")
    except ValueError:
        return None


def _decimal(value: Optional[str]) -> Optional[Decimal]:
    try:
        return Decimal(value) if value else None
    except InvalidOperation:
        return None


def _store(
    session: Session,
    layout: OneDriveLayout,
    path: Path,
    texts: list[tuple[str, str]],
    stat: Optional[os.stat_result] = None,
    sha256: Optional[str] = None,
) -> ScannedDocument:
    rel = path.relative_to(layout.onedrive)
    parts = list(rel.parts)
    key = str(rel).replace("\\", "/")
    parsed = parse_filename(path.name)
    joined = "\n".join(t for t, _ in texts)
    ocr_date, _ = find_date_and_time(joined) if joined.strip() else (None, None)
    sources = {s for _, s in texts if s != "none"}
    stat = stat or os.stat(path)

    doc = session.scalar(select(ScannedDocument).where(ScannedDocument.path == key))
    if doc is None:
        doc = ScannedDocument(path=key)
        session.add(doc)
    doc.folder = str(rel.parent).replace("\\", "/")
    doc.person = parts[1] if parts[0].lower() == "people" and len(parts) > 2 else None
    doc.kind = _kind_for(parts)
    doc.doc_date = _parse_date(parsed.date)
    doc.doc_time = parsed.time.replace("-", ":") if parsed.time else None
    doc.merchant = parsed.slug or path.stem
    doc.page_count = len(texts)
    doc.file_size = stat.st_size
    doc.file_modified = datetime.fromtimestamp(stat.st_mtime)
    doc.sha256 = sha256 or doc.sha256
    doc.ocr_date = ocr_date
    doc.ocr_total = _decimal(find_total(joined))
    doc.text_source = (sources.pop() if len(sources) == 1 else "mixed") if sources else "none"
    doc.indexed_at = datetime.now()
    doc.pages = [ScannedPage(page_number=n, text=t, source=s) for n, (t, s) in enumerate(texts, start=1)]
    return doc


def file_sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with open(path, "rb") as handle:
        for chunk in iter(lambda: handle.read(1 << 20), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _key(layout: OneDriveLayout, path: Path) -> str:
    return str(Path(path).relative_to(layout.onedrive)).replace("\\", "/")


@dataclass
class IndexReport:
    indexed: list[str] = field(default_factory=list)  # text read (OCR or text layer)
    reused: list[str] = field(default_factory=list)  # same content already known: no OCR
    moved: int = 0  # renamed or moved files, carried over
    unchanged: int = 0
    skipped_online_only: int = 0
    removed: int = 0
    errors: list[str] = field(default_factory=list)


def index_documents(
    engine: Engine,
    layout: OneDriveLayout,
    ocr: Optional[Callable],
    roots: Optional[Iterable[Path]] = None,
    include_online_only: bool = False,
    known_text: Optional[dict[str, list[str]]] = None,
    progress: Optional[Callable[[int, int], None]] = None,
    commit_every: int = 1,  # short write locks: other commands can use the database meanwhile
    files: Optional[Iterable[Path]] = None,
) -> IndexReport:
    """Index every PDF under ``roots`` (or exactly ``files``), reading each file's text only once.

    A file is read again only if its content changes. Unchanged content found
    under a new name or folder (renamed, moved, or re-dated with set-dates)
    keeps its stored text. ``known_text`` imports OCR done elsewhere
    (path -> page texts).
    """
    from pypdf import PdfReader

    from .filedates import is_online_only

    ensure_schema(engine)
    report = IndexReport()
    explicit = files is not None
    roots = list(roots or default_roots(layout))
    files = sorted(set(files) if explicit else {p for root in roots for p in Path(root).rglob("*.pdf")})
    with Session(engine) as session:
        by_path = {d.path: d for d in session.scalars(select(ScannedDocument))}
        by_hash = {d.sha256: d for d in by_path.values() if d.sha256}
        for index, path in enumerate(files, start=1):
            if progress:
                progress(index, len(files))
            key = _key(layout, path)
            try:
                stat = os.stat(path)
                doc = by_path.get(key)
                if (doc and doc.file_size == stat.st_size
                        and doc.file_modified == datetime.fromtimestamp(stat.st_mtime)):
                    report.unchanged += 1
                    continue
                if not include_online_only and is_online_only(path):
                    report.skipped_online_only += 1
                    continue
                digest = file_sha256(path)
                if doc and doc.sha256 == digest:
                    doc.file_modified = datetime.fromtimestamp(stat.st_mtime)  # only times changed
                    report.unchanged += 1
                    continue
                twin = by_hash.get(digest)
                if twin is not None and twin.path != key and not (layout.onedrive / twin.path).exists():
                    # Renamed or moved: carry the row over to the new path.
                    by_path.pop(twin.path, None)
                    texts = [(p.text, p.source) for p in twin.pages]
                    session.delete(twin)
                    session.flush()
                    report.moved += 1
                elif twin is not None:
                    texts = [(p.text, p.source) for p in twin.pages]  # a copy elsewhere
                    report.reused.append(key)
                else:
                    pages = len(PdfReader(str(path)).pages)
                    cached = (known_text or {}).get(key)
                    if cached is not None and len(cached) == pages:
                        texts = [(t, "ocr:imported" if t.strip() else "none") for t in cached]
                        report.reused.append(key)
                    else:
                        texts = [page_text_with_source(PageRef(path, n), ocr) for n in range(1, pages + 1)]
                        report.indexed.append(key)
                stored = _store(session, layout, path, texts, stat, digest)
                by_path[key] = by_hash[digest] = stored
            except Exception as exc:  # a damaged PDF should not stop the run
                report.errors.append(f"{key}: {exc}")
            if index % commit_every == 0:
                session.commit()
        if not explicit:
            # Forget documents whose files are gone and were not moved.
            scanned = [Path(r).resolve() for r in roots]
            for key, doc in list(by_path.items()):
                full = layout.onedrive / key
                if not full.exists() and any(full.resolve().is_relative_to(r) for r in scanned):
                    session.delete(doc)
                    report.removed += 1
        session.commit()
    return report


def stored_texts(engine: Engine, layout: OneDriveLayout, paths: Iterable[Path]) -> dict[str, list[str]]:
    """Page texts already in the database, keyed by path relative to OneDrive."""
    ensure_schema(engine)
    keys = [_key(layout, p) for p in paths]
    with Session(engine) as session:
        docs = session.scalars(select(ScannedDocument).where(ScannedDocument.path.in_(keys)))
        return {d.path: [p.text for p in d.pages] for d in docs}


def all_stored_texts(engine: Engine) -> dict[str, list[str]]:
    ensure_schema(engine)
    with Session(engine) as session:
        return {d.path: [p.text for p in d.pages] for d in session.scalars(select(ScannedDocument))}


def index_filed_document(engine: Engine, layout: OneDriveLayout, path: Path, page_texts: list[str]) -> None:
    """Record a document just filed by ``apply``, using text read at plan time."""
    ensure_schema(engine)
    with Session(engine) as session:
        texts = [(t, "ocr:plan" if t.strip() else "none") for t in page_texts]
        _store(session, layout, path, texts, sha256=file_sha256(path))
        session.commit()


def forget_documents(engine: Engine, paths: Iterable[str]) -> int:
    """Drop index rows for paths relative to OneDrive (after ``undo``)."""
    ensure_schema(engine)
    keys = [str(p).replace("\\", "/") for p in paths]
    with Session(engine) as session:
        docs = list(session.scalars(select(ScannedDocument).where(ScannedDocument.path.in_(keys))))
        for doc in docs:
            session.delete(doc)
        session.commit()
    return len(docs)


@dataclass
class SearchHit:
    path: str
    page: int
    snippet: str
    doc_date: Optional[date]


def search(engine: Engine, query: str, limit: int = 20) -> list[SearchHit]:
    """FTS5 query: words, "exact phrases", OR, prefix*."""
    ensure_schema(engine)
    sql = text(
        """
        SELECT d.path, p.page_number, snippet(scanned_pages_fts, 0, '[', ']', ' ... ', 12), d.doc_date
        FROM scanned_pages_fts
        JOIN scanned_pages p ON p.id = scanned_pages_fts.rowid
        JOIN scanned_documents d ON d.id = p.document_id
        WHERE scanned_pages_fts MATCH :query
        ORDER BY rank
        LIMIT :limit
        """
    )
    with engine.connect() as connection:
        rows = connection.execute(sql, {"query": query, "limit": limit}).all()
    return [SearchHit(r[0], r[1], r[2], _parse_date(r[3]) if isinstance(r[3], str) else r[3]) for r in rows]


def export_xlsx(engine: Engine, out: Path) -> tuple[int, int]:
    """Two sheets: documents (one row per PDF) and pages (one row per page)."""
    from openpyxl import Workbook
    from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
    from openpyxl.styles import Alignment, Font

    def cell_text(value: str) -> str:
        # OCR output can contain control characters Excel refuses.
        return ILLEGAL_CHARACTERS_RE.sub("", value)[:EXCEL_CELL_LIMIT]

    ensure_schema(engine)
    workbook = Workbook()
    docs_sheet = workbook.active
    docs_sheet.title = "documents"
    doc_columns = ["path", "folder", "person", "kind", "doc_date", "doc_time", "merchant", "page_count",
                   "ocr_date", "ocr_total", "text_source", "indexed_at", "text"]
    docs_sheet.append(doc_columns)
    pages_sheet = workbook.create_sheet("pages")
    pages_sheet.append(["path", "page", "source", "text"])
    n_docs = n_pages = 0
    with Session(engine) as session:
        for doc in session.scalars(select(ScannedDocument).order_by(ScannedDocument.path)):
            full = "\n\n".join(p.text for p in doc.pages)
            docs_sheet.append([
                doc.path, doc.folder, doc.person, doc.kind, doc.doc_date, doc.doc_time, doc.merchant,
                doc.page_count, doc.ocr_date, float(doc.ocr_total) if doc.ocr_total is not None else None,
                doc.text_source, doc.indexed_at, cell_text(full),
            ])
            n_docs += 1
            for page in doc.pages:
                pages_sheet.append([doc.path, page.page_number, page.source, cell_text(page.text)])
                n_pages += 1
    for sheet, widths in ((docs_sheet, [60, 35, 18, 12, 12, 9, 30, 8, 12, 10, 10, 18, 80]),
                          (pages_sheet, [60, 6, 10, 100])):
        for cell in sheet[1]:
            cell.font = Font(bold=True)
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for column, width in zip("ABCDEFGHIJKLM", widths):
            sheet.column_dimensions[column].width = width
        text_column = sheet.max_column
        for row in sheet.iter_rows(min_row=2, min_col=text_column, max_col=text_column):
            row[0].alignment = Alignment(wrap_text=False, vertical="top")
    out.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(out)
    return n_docs, n_pages
